"""Build site/event_cost_lines.json, the proof behind each event's cost.

Reads the month-end event cost workbook (CHIIRP Event Costs.xlsx, built by
the runbook in EVENT_COSTS_MONTHLY.md: Event Hub budgets and attendees, and
Ramp card charges matched to events) and writes one entry per event: the fee
as the Event Hub records it, every matched card charge with its owner, date,
merchant, label and why it was matched, and the charges reviewed and not
counted. The site shows it on each event page under "Proof of cost".

    py -3 scripts/build_event_cost_lines.py "C:/Users/josha/OneDrive/Desktop/CHIIRP Event Costs.xlsx"

Events before the Event Hub (Aug 2025 to Aug 2026) come from a second
workbook, CHIIRP Event Costs 2025-26 (QuickBooks).xlsx, read when it sits
beside the first: every 6100-07 bill as the fee, and every card charge on any
card (Ramp, Divvy, the company Chase card) from QuickBooks' travel accounts,
with the merchant city from the Ramp export, matched by the same rules.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "site" / "event_cost_lines.json"

# The workbook names events as the Event Hub does; the site names them as the
# Event Costs tab does.
LABELS = {
    "Grosso University Extreme Leadership Workshop": "Grosso University 2026",
    "Home Service Hoorah": "Home Service Hoorah 2026",
    "Nexstar Super Meeting": "Nexstar Super Meeting 2026",
    "The Dirty Dozen Retreat": "Dirty Dozen Retreat 2026",
    "CCN ELEVATE 2026": "CCN Elevate 2026",
    "Service Titan Pantheon": "Pantheon 2026",
    "CertainPath Connect/Partner Jam": "CertainPath Connect 2026",
    "Qualified Remodeler Top 500": "Qualified Remodeler Top 500 2026",
    "HSF Freedom 2026": "Home Service Freedom 2026",
    "Build it to Sell it": "Build it to Sell it 2026",
    "Table of Titans (Tony Hoty Mastermind)": "Table of Titans 2027",
    "EGIA EPIC 2027": "EGIA Epic 2027",
}

# Attendees as the Event Hub lists them on 13 Oct 2026 (Josh), for events the
# workbook's Ramp Travel Summary does not cover yet (upcoming). The workbook
# wins wherever it names attendees.
HUB_ATTENDEES = {
    "CertainPath Connect/Partner Jam": ["Ryan Peterson", "Justin Judd"],
    "Qualified Remodeler Top 500": ["Justin Judd", "Ryan Fenn", "Ky Decker", "Ryan Peterson"],
    "HSF Freedom 2026": ["Colton Lingmann", "Blake Fenn", "Taylor Roberts", "Ryan Fenn"],
    "Table of Titans (Tony Hoty Mastermind)": ["Justin Judd (not registered)", "Ryan Fenn (not registered)"],
}
# Judgement calls about attendees made after the cost run.
HUB_NOTES = {
    "Home Service Hoorah": "The Event Hub no longer lists Ryan Fenn (13 Oct 2026), but his card puts him in Gulf Breeze on the "
                           "event dates (a Delta flight, a Hilton Gulf Breeze stay, a rental car), so his charges stay counted.",
}


def iso(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v) if v not in (None, "") else None


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        if not r or r[0] in (None, "") or str(r[0]).startswith(("GRAND TOTAL", "Source:", "Sources:")):
            continue
        yield dict(zip(head, r))


def older(path):
    """Events before the Event Hub, from the QuickBooks workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for r in rows(wb["Events"]):
        out[r["Event"]] = {
            "label": r["Event"], "hubName": None, "start": iso(r.get("Start date")), "end": iso(r.get("End date")),
            "location": r.get("Location"), "venue": None, "basis": "quickbooks",
            "fee": {"amount": num(r.get("Fee in QuickBooks ($)")), "paymentStatus": None, "what": None,
                    "source": "QuickBooks, 6100-07 Tradeshows & Events"},
            "feeLines": [], "hubTravel": None, "flags": None,
            "attendees": [a.strip() for a in str(r.get("Attendees") or "").split(",") if a.strip()],
            "notes": r.get("Notes"), "lines": [], "notCounted": [],
        }
    for r in rows(wb["Fee Lines"]):
        e = out.get(r["Event"])
        if e:
            e["feeLines"].append({"date": iso(r.get("Date")), "vendor": r.get("Vendor"), "what": r.get("What"),
                                  "number": r.get("Number"), "type": r.get("Type"), "amount": num(r.get("Amount ($)")),
                                  "expensed": r.get("Expensed in"), "confidence": r.get("Confidence"), "note": r.get("Note")})
    for r in rows(wb["Card Detail"]):
        e = out.get(r["Event"])
        if e:
            e["lines"].append({"date": iso(r.get("Date")), "owner": r.get("Cardholder"), "merchant": r.get("Merchant"),
                               "rampCategory": None, "city": r.get("Merchant city"), "amount": num(r.get("Amount ($)")),
                               "label": r.get("Category"), "confidence": r.get("Confidence"), "why": r.get("Why it was matched"),
                               "card": r.get("Card"), "source": f"QuickBooks, {r.get('Card')} card"})
    for r in rows(wb["Needs Review"]):
        e = out.get(r["Nearest event"]) if r.get("Nearest event") else None
        if e:
            e["notCounted"].append({"date": iso(r.get("Date")), "owner": r.get("Cardholder"), "merchant": r.get("Merchant"),
                                    "city": r.get("Merchant city"), "amount": num(r.get("Amount ($)")),
                                    "why": r.get("Why it is not counted")})
    return out


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    events = {}
    for r in rows(wb["Event Costs"]):
        name = str(r["Event"]).strip()
        if name not in LABELS:
            continue
        events[name] = {
            "label": LABELS[name], "hubName": name, "start": iso(r.get("Start date")), "end": iso(r.get("End date")),
            "location": r.get("Location"), "venue": r.get("Venue"),
            "fee": {"amount": num(r.get("Sponsorship / Booth ($)")), "paymentStatus": r.get("Payment status"),
                    "what": r.get("Cost notes (from app)"), "source": "Event Hub, Budget & payment"},
            "hubTravel": {"flights": num(r.get("Flights ($)")), "hotels": num(r.get("Hotels ($)")),
                          "rentalCars": num(r.get("Rental cars ($)"))},
            "flags": r.get("Notes / flags"), "attendees": [], "notes": None, "lines": [], "notCounted": [],
        }
    for r in rows(wb["Ramp Travel Summary"]):
        e = events.get(str(r["Event"]).strip())
        if e:
            e["attendees"] = [a.strip() for a in str(r.get("Attendees (Event Hub)") or "").split(",") if a.strip()]
            e["notes"] = r.get("Notes")
    for r in rows(wb["Ramp Detail"]):
        e = events.get(str(r["Event"]).strip())
        if not e:
            continue
        e["lines"].append({"date": iso(r.get("Date")), "owner": r.get("Cardholder"), "merchant": r.get("Merchant"),
                           "rampCategory": r.get("Ramp category"), "city": r.get("Merchant city"),
                           "amount": num(r.get("Amount ($)")), "label": r.get("Category"),
                           "confidence": r.get("Confidence"), "why": r.get("Why it was matched / note"),
                           "source": "Ramp card"})
    # Reviewed and not counted: the review tab names the event in its reason.
    for r in rows(wb["Ramp Needs Review"]):
        why = str(r.get("Why it is not counted") or "")
        for name, e in events.items():
            if name in why:
                e["notCounted"].append({"date": iso(r.get("Date")), "owner": r.get("Cardholder"), "merchant": r.get("Merchant"),
                                        "city": r.get("Merchant city"), "amount": num(r.get("Amount ($)")), "why": why})
    qb = Path(path).with_name("CHIIRP Event Costs 2025-26 (QuickBooks).xlsx")
    if qb.exists():
        for label, e in older(qb).items():
            if not any(x["label"] == label for x in events.values()):
                events[f"qb:{label}"] = e
    for name, e in events.items():
        if not e["attendees"] and name in HUB_ATTENDEES:
            e["attendees"] = HUB_ATTENDEES[name]
        if name in HUB_NOTES:
            e["notes"] = " ".join(x for x in (e["notes"], HUB_NOTES[name]) if x)
    for e in events.values():
        e.setdefault("basis", "event hub")
        e["lines"].sort(key=lambda x: (str(x["owner"]), str(x["date"])))
        e["travelTotal"] = round(sum(x["amount"] or 0 for x in e["lines"]), 2)
    out = {"source": Path(path).name, "olderSource": qb.name if qb.exists() else None, "built": date.today().isoformat(),
           "note": "Event Hub fees and Ramp card charges matched to each event by the month-end cost run; events before "
                   "the Event Hub from QuickBooks bills and card charges.",
           "events": sorted(events.values(), key=lambda e: e["start"] or "")}
    OUT.write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8", newline="\n")
    print(f"{OUT}: {len(out['events'])} events, {sum(len(e['lines']) for e in out['events'])} charges, "
          f"{sum(len(e['notCounted']) for e in out['events'])} reviewed and not counted")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\josha\OneDrive\Desktop\CHIIRP Event Costs.xlsx")
